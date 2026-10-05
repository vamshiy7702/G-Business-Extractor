import asyncio
import io
import json
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from leadx.adapters.osm_overpass import parse_element
from leadx.api import create_app
from leadx.db import Store


def mk(i, phone=True, rating=None):
    tags = {"name": f"Clinic {i}", "amenity": "dentist", "website": f"https://c{i}.example"}
    if phone:
        tags["phone"] = f"+91 {i}"
    return parse_element({"type": "node", "id": i, "lat": 12.0, "lon": 77.0, "tags": tags})


class FakeAdapter:
    def __init__(self, delay=0.0):
        self.delay, self.calls = delay, 0

    async def search(self, keyword, tile, limit):
        self.calls += 1
        await asyncio.sleep(self.delay)
        b = self.calls * 10
        return [mk(b), mk(b + 1, phone=False)]

    def estimate_cost(self, **kwargs):
        from leadx.adapters.base import Quote
        return Quote("fake", 0, 100000, 0, 0, 100000)


async def fake_geocoder(client, city, state="", country="", postal_code=""):
    if city == "Nowhere":
        raise ValueError("Could not find location: Nowhere")
    if city == "Huge":
        return (0.0, 0.0, 100.0, 1000.0)
    return (12.0, 77.0, 12.2, 77.2)


BODY = {"keyword": "dentist", "location": {"city": "Bangalore", "country": "IN"},
        "options": {"max_places": 100, "tile_deg": 0.1}}


def make_client(tmp_path, delay=0.0):
    app = create_app(str(tmp_path / "api.db"), adapter_factory=lambda source: FakeAdapter(delay), geocoder=fake_geocoder)
    return TestClient(app)


def wait_for(c, jid, statuses, timeout=10):
    end = time.time() + timeout
    while time.time() < end:
        j = c.get(f"/api/jobs/{jid}").json()
        if j["status"] in statuses:
            return j
        time.sleep(0.05)
    raise AssertionError(f"job {jid} never reached {statuses}: {j}")


def test_health_and_unknown_job(tmp_path):
    with make_client(tmp_path) as c:
        assert c.get("/api/health").json()["status"] == "ok"
        assert c.get("/api/jobs/999").status_code == 404
        assert c.get("/api/jobs/999/places").status_code == 404
        assert c.get("/api/settings").status_code == 200
        assert c.get("/api/quota").status_code == 200


def test_create_run_list_filter_page_export(tmp_path):
    with make_client(tmp_path) as c:
        r = c.post("/api/jobs", json=BODY)
        assert r.status_code == 201
        jid = r.json()["id"]
        job = wait_for(c, jid, {"done"})
        assert job["places_found"] == 8 and job["progress"] == 100.0 and job["tiles_total"] == 4
        allp = c.get(f"/api/jobs/{jid}/places").json()
        assert allp["total"] == 8
        phones = c.get(f"/api/jobs/{jid}/places", params={"has_phone": True}).json()
        assert phones["total"] == 4 and all(p["phone"] for p in phones["items"])
        page = c.get(f"/api/jobs/{jid}/places", params={"offset": 6, "limit": 5}).json()
        assert len(page["items"]) == 2 and page["total"] == 8
        assert c.get(f"/api/jobs/{jid}/places", params={"q": "Clinic 10"}).json()["total"] == 1
        csv_r = c.get(f"/api/jobs/{jid}/export", params={"format": "csv", "has_phone": True})
        lines = csv_r.content.decode("utf-8-sig").strip().splitlines()
        assert "attachment" in csv_r.headers["content-disposition"] and len(lines) == 5
        x = c.get(f"/api/jobs/{jid}/export", params={"format": "xlsx"})
        ws = load_workbook(io.BytesIO(x.content)).active
        assert ws.max_row == 9 and ws["C1"].value == "name"
        assert c.get(f"/api/jobs/{jid}/export", params={"format": "pdf"}).status_code == 422


def test_validation_and_errors(tmp_path):
    with make_client(tmp_path) as c:
        assert c.post("/api/jobs", json={**BODY, "keyword": ""}).status_code == 422
        assert c.post("/api/jobs", json={**BODY, "options": {"max_places": 0}}).status_code == 422
        r = c.post("/api/jobs", json={**BODY, "location": {"city": "Nowhere", "country": "IN"}})
        assert r.status_code == 404 and "Nowhere" in r.json()["detail"]
        r = c.post("/api/jobs", json={**BODY, "location": {"city": "Huge", "country": "IN"}, "options": {"tile_deg": 0.01}})
        assert r.status_code == 400 and "too large" in r.json()["detail"]


def test_all_location_scopes_and_global_guard(tmp_path):
    seen = []

    async def capture_geocoder(client, city, state="", country="", postal_code=""):
        seen.append((city, state, country, postal_code))
        return (12.0, 77.0, 12.2, 77.2)

    app = create_app(str(tmp_path / "scope.db"), adapter_factory=lambda source: FakeAdapter(), geocoder=capture_geocoder)
    with TestClient(app) as c:
        body = {**BODY, "location": {"city": "All cities", "state": "Telangana", "country": "India", "postal_code": "All zip codes"}}
        r = c.post("/api/jobs", json=body)
        assert r.status_code == 201
        assert seen[-1] == ("All cities", "Telangana", "India", "All zip codes")
        jid = r.json()["id"]
        assert wait_for(c, jid, {"done"})["status"] == "done"

    async def guarded_geocoder(client, city, state="", country="", postal_code=""):
        if country.lower() == "all countries":
            raise ValueError("All countries cannot be geocoded as one bounded OSM search area.")
        return (12.0, 77.0, 12.2, 77.2)

    guarded_app = create_app(str(tmp_path / "global.db"), adapter_factory=lambda source: FakeAdapter(), geocoder=guarded_geocoder)
    with TestClient(guarded_app) as c:
        r = c.post("/api/jobs", json={**BODY, "location": {"city": "All cities", "state": "All states", "country": "All countries", "postal_code": "All zip codes"}})
        assert r.status_code == 404
        assert "All countries" in r.json()["detail"]


def test_bulk_delete_selected_and_all(tmp_path):
    with make_client(tmp_path) as c:
        first = c.post("/api/jobs", json=BODY).json()["id"]
        second = c.post("/api/jobs", json={**BODY, "keyword": "restaurant"}).json()["id"]
        assert wait_for(c, first, {"done"})["status"] == "done"
        assert wait_for(c, second, {"done"})["status"] == "done"

        r = c.post("/api/jobs/bulk-delete", json={"ids": [first]})
        assert r.status_code == 200 and r.json()["deleted"] == [first]
        assert c.get(f"/api/jobs/{first}").status_code == 404
        assert c.get(f"/api/jobs/{second}").status_code == 200

        r = c.post("/api/jobs/bulk-delete", json={"ids": [second, 999999]})
        assert r.status_code == 200
        assert r.json()["deleted"] == [second]
        assert r.json()["missing"] == [999999]
        assert c.get("/api/jobs").json() == []


def test_cancel_then_resume_and_delete(tmp_path):
    with make_client(tmp_path, delay=0.15) as c:
        jid = c.post("/api/jobs", json=BODY).json()["id"]
        end = time.time() + 5
        while c.get(f"/api/jobs/{jid}").json()["tiles_done"] < 1 and time.time() < end:
            time.sleep(0.05)
        c.post(f"/api/jobs/{jid}/cancel")
        paused = wait_for(c, jid, {"paused"})
        assert 1 <= paused["tiles_done"] < 4
        assert c.post(f"/api/jobs/{jid}/resume").status_code == 200
        assert wait_for(c, jid, {"done"})["tiles_done"] == 4
        assert c.post(f"/api/jobs/{jid}/resume").status_code == 409
        assert c.delete(f"/api/jobs/{jid}").status_code == 204
        assert c.get(f"/api/jobs/{jid}").status_code == 404


def test_sse_stream_ends_with_all_places(tmp_path):
    with make_client(tmp_path) as c:
        jid = c.post("/api/jobs", json=BODY).json()["id"]
        seen, end_status, event = [], None, None
        with c.stream("GET", f"/api/jobs/{jid}/stream") as resp:
            assert resp.headers["content-type"].startswith("text/event-stream")
            for line in resp.iter_lines():
                if line.startswith("event:"):
                    event = line.split(":", 1)[1].strip()
                elif line.startswith("data:"):
                    data = json.loads(line.split(":", 1)[1])
                    if event == "progress":
                        seen += [p["id"] for p in data["new_places"]]
                    elif event == "end":
                        end_status = data["status"]
                if end_status:
                    break
        assert end_status == "done"
        assert len(seen) == len(set(seen)) == 8


def test_stale_running_jobs_become_paused_on_startup(tmp_path):
    dbp = str(tmp_path / "api.db")
    s = Store(dbp)
    jid = s.create_job(keyword="x", bbox=(0, 0, 1, 1), tiles_total=4, tile_deg=0.5,
                       max_places=5, per_tile=5)
    s.set_status(jid, "running")
    s.close()
    app = create_app(dbp, adapter_factory=lambda source: FakeAdapter(), geocoder=fake_geocoder)
    with TestClient(app) as c:
        assert c.get(f"/api/jobs/{jid}").json()["status"] == "paused"


def test_delete_selected_and_all_places(tmp_path):
    with make_client(tmp_path) as c:
        jid = c.post("/api/jobs", json=BODY).json()["id"]
        assert wait_for(c, jid, {"done"})["places_found"] == 8
        places = c.get(f"/api/jobs/{jid}/places").json()["items"]
        ids = [places[0]["id"], places[1]["id"]]
        r = c.post(f"/api/jobs/{jid}/places/delete", json={"ids": ids})
        assert r.status_code == 200 and r.json()["count"] == 2
        assert c.get(f"/api/jobs/{jid}/places").json()["total"] == 6
        r = c.post(f"/api/jobs/{jid}/places/delete-all")
        assert r.status_code == 200 and r.json()["count"] == 6
        assert c.get(f"/api/jobs/{jid}/places").json()["total"] == 0
